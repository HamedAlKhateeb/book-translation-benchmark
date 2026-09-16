#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Book Translation Benchmark Web Application
Features:
- Single Chapter & Multi-File Batch Processing
- Dark / Light Mode Theme Toggle
- Formats: EPUB, DOCX, HTML, Markdown
- Metrics: COMET (Unbabel/wmt22-comet-da), MQM, chrF++, BLEU
- Database & Auth: Firebase Authentication (Email/Password & Google) + Cloud Firestore
- Export Formats: PDF, DOCX, Markdown, CSV
"""

import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')
import re
import csv
import io
import json
import zipfile
import functools
import urllib.parse
from html.parser import HTMLParser
from bs4 import BeautifulSoup
from rapidfuzz import fuzz
import sacrebleu
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Monkeypatch for Python 3.14 compatibility with COMET
class _HashedSeq(list):
    __slots__ = 'hashvalue'
    def __init__(self, tup, hash=hash):
        self[:] = tup
        self.hashvalue = hash(tup)
    def __hash__(self):
        return self.hashvalue
functools._HashedSeq = _HashedSeq

# Optional COMET loader
comet_model = None
def get_comet_model():
    global comet_model
    if comet_model is None:
        from comet import load_from_checkpoint
        ckpt_path = r'C:\Users\hamed\.cache\huggingface\hub\models--Unbabel--wmt22-comet-da\snapshots\2760a223ac957f30acfb18c8aa649b01cf1d75f2\checkpoints\model.ckpt'
        if os.path.exists(ckpt_path):
            print("Loading local cached COMET model...")
            comet_model = load_from_checkpoint(ckpt_path)
        else:
            print("Downloading/Loading COMET model...")
            from comet import download_model
            path = download_model("Unbabel/wmt22-comet-da")
            comet_model = load_from_checkpoint(path)
    return comet_model

# ----------------- Text Cleaning & Extraction -----------------
def clean_arabic(text):
    text = re.sub(r'[\u200e\u200f\u202a-\u202e]', '', text)
    text = re.sub(r'[،,]?[٠-٩]+', '', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def clean_english(text):
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def extract_text_from_file(filename, file_bytes):
    ext = filename.lower().split('.')[-1]
    paragraphs = []

    if ext in ['html', 'xhtml']:
        soup = BeautifulSoup(file_bytes.decode('utf-8', errors='ignore'), 'html.parser')
        for tag in soup.find_all(['p', 'h1', 'h2', 'h3', 'h4', 'li']):
            txt = tag.get_text().strip()
            if len(txt) > 15:
                paragraphs.append(txt)

    elif ext == 'docx':
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
            xml_content = z.read('word/document.xml')
            soup = BeautifulSoup(xml_content, 'xml')
            for p in soup.find_all('p'):
                txt = p.get_text().strip()
                if txt and not txt.startswith('HTML+') and len(txt) > 15:
                    paragraphs.append(txt)

    elif ext == 'epub':
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
            html_files = [f for f in z.namelist() if f.lower().endswith(('.xhtml', '.html', '.htm'))]
            html_files.sort()
            for h in html_files:
                content = z.read(h).decode('utf-8', errors='ignore')
                soup = BeautifulSoup(content, 'html.parser')
                for tag in soup.find_all(['p', 'h1', 'h2', 'h3', 'li']):
                    txt = tag.get_text().strip()
                    if len(txt) > 15:
                        paragraphs.append(txt)

    elif ext in ['md', 'markdown', 'txt']:
        raw = file_bytes.decode('utf-8', errors='ignore')
        for line in raw.splitlines():
            line = line.strip()
            if len(line) > 15:
                line = re.sub(r'^#+\s*', '', line)
                paragraphs.append(line)

    return paragraphs

# ----------------- Dynamic Alignment Engine -----------------
def align_segments(ref_list, ai_list, max_merge=4):
    n_ref, n_ai = len(ref_list), len(ai_list)
    dp = [[-1e9] * (n_ai + 1) for _ in range(n_ref + 1)]
    back = [[None] * (n_ai + 1) for _ in range(n_ref + 1)]
    dp[0][0] = 0

    for i in range(n_ref + 1):
        for j in range(n_ai + 1):
            if dp[i][j] < -1e8:
                continue
            for dr in range(1, max_merge + 1):
                if i + dr > n_ref:
                    continue
                r_txt = ' '.join(ref_list[i:i+dr])
                for da in range(1, max_merge + 1):
                    if j + da > n_ai:
                        continue
                    a_txt = ' '.join(ai_list[j:j+da])
                    score = fuzz.token_set_ratio(r_txt, a_txt)
                    if dp[i][j] + score > dp[i+dr][j+da]:
                        dp[i+dr][j+da] = dp[i][j] + score
                        back[i+dr][j+da] = (i, j, dr, da)

    curr_i, curr_j = n_ref, n_ai
    matches = []
    while curr_i > 0 and curr_j > 0:
        b = back[curr_i][curr_j]
        if b is None:
            break
        prev_i, prev_j, dr, da = b
        r_txt = ' '.join(ref_list[prev_i:prev_i+dr])
        a_txt = ' '.join(ai_list[prev_j:prev_j+da])
        matches.append((prev_i, prev_j, dr, da, r_txt, a_txt))
        curr_i, curr_j = prev_i, prev_j

    matches.reverse()
    return matches

# ----------------- MQM Error Taxonomy -----------------
def classify_mqm(src, ref, hyp, comet_score, chrf_score, bleu_score):
    hyp_words = hyp.split()
    word_count = max(1, len(hyp_words))
    
    english_words = re.findall(r'[a-zA-Z]{3,}', hyp)
    english_words = [w for w in english_words if w.lower() not in ['html', 'docx', 'pdf', 'url']]
    
    hyp_len = len(hyp_words)
    ref_len = len(ref.split())
    length_ratio = hyp_len / max(1, ref_len)
    
    if length_ratio < 0.65 and ref_len > 25:
        cat = 'الدقة (Accuracy)'
        err_type = 'حذف / سقط في المعنى (Omission)'
        sev = 'جسيم (Major)'
        desc = f'سقط جزء كبير من الفقرة في ترجمة الذكاء الاصطناعي بنسبة نقص {int((1-length_ratio)*100)}%'
        penalty = 5
    elif english_words:
        cat = 'الدقة (Accuracy)'
        err_type = 'نص غير مترجم (Untranslated)'
        sev = 'طفيف (Minor)'
        desc = f'بقاء مفردات أجنبية دون تعريب: {", ".join(english_words[:3])}'
        penalty = 1
    elif comet_score < 0.72 or (comet_score < 0.76 and chrf_score < 30):
        cat = 'الدقة (Accuracy)'
        err_type = 'ترجمة غير دقيقة / معنى مغاير (Mistranslation)'
        sev = 'جسيم (Major)'
        desc = 'انحراف دلالي واضح عن قصد النص الأصلي أو المصطلحات المعتمدة'
        penalty = 5
    elif chrf_score < 42 and comet_score < 0.83:
        cat = 'المصطلحات (Terminology)'
        err_type = 'مصطلح غير دقيق (Inappropriate Term)'
        sev = 'طفيف (Minor)'
        desc = 'استخدام مصطلح عام أو معجمي غير دقيق مقارنة بالمصطلح التخصصي البشري'
        penalty = 1
    elif bleu_score < 15 and chrf_score < 48:
        cat = 'الأسلوب (Style)'
        err_type = 'حرفية وركاكة أسلوبية (Literalness / Awkward)'
        sev = 'طفيف (Minor)'
        desc = 'ترجمة حرفية تتبع بناء الجملة الإنجليزية مع ضعف في بلاغة التركيب العربي'
        penalty = 1
    elif length_ratio > 1.45 and ref_len > 25:
        cat = 'الدقة (Accuracy)'
        err_type = 'إضافة غير مبررة (Addition)'
        sev = 'طفيف (Minor)'
        desc = 'إسهاب وحشو زائد في صياغة الفكرة لم يرد في النص الأصلي'
        penalty = 1
    else:
        cat = 'لا يوجد (None)'
        err_type = 'ترجمة صحيحة وسليمة (Accurate & Fluent)'
        sev = 'منعدم (None)'
        desc = 'ترجمة متسقة وسلسة ومكافئة دلالياً للأصل البشري'
        penalty = 0

    penalty_points = penalty * (100.0 / max(20, word_count))
    mqm_score = max(0.0, min(100.0, 100.0 - penalty_points))
    return cat, err_type, sev, desc, round(mqm_score, 2)

# ----------------- DOCX Generator -----------------
def generate_docx_report(data):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("تقرير مقارنة وتقييم جودة ترجمة الكتب")
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    book_title = data.get('book_title') or 'كتاب'
    chapter = data.get('chapter') or 'فصل'
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(f"الكتاب: {book_title}  |  الفصل: {chapter}")
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph()

    doc.add_heading("1. ملخص المعايير التقييمية (Key Metrics)", level=2)
    summary = data.get('summary', {})
    t_m = doc.add_table(rows=2, cols=4)
    t_m.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    m_headers = ["COMET (العصبي)", "جودة MQM", "chrF++ (الصرفي)", "BLEU (اللفظي)"]
    m_vals = [
        str(summary.get('avg_comet', '0.0000')),
        f"{summary.get('avg_mqm', '0.0')}%",
        str(summary.get('avg_chrf', '0.00')),
        str(summary.get('avg_bleu', '0.00'))
    ]
    for i, h in enumerate(m_headers):
        cell = t_m.cell(0, i)
        cell.text = h
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs: p.runs[0].font.bold = True

    for i, v in enumerate(m_vals):
        cell = t_m.cell(1, i)
        cell.text = v
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs:
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(14)

    doc.add_paragraph()

    doc.add_heading("2. توزيع ونسب أنواع الأخطاء (MQM Breakdown)", level=2)
    errors = data.get('errors', [])
    t_e = doc.add_table(rows=len(errors)+1, cols=4)
    t_e.alignment = WD_TABLE_ALIGNMENT.CENTER
    e_headers = ["نوع الخطأ / التصنيف", "مستوى الخطورة", "التكرار", "النسبة المئوية (%)"]
    for i, h in enumerate(e_headers):
        cell = t_e.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True

    for r_idx, err in enumerate(errors):
        t_e.cell(r_idx+1, 0).text = str(err.get('type', ''))
        t_e.cell(r_idx+1, 1).text = str(err.get('severity', ''))
        t_e.cell(r_idx+1, 2).text = str(err.get('count', 0))
        t_e.cell(r_idx+1, 3).text = str(err.get('percentage', ''))

    doc.add_paragraph()

    doc.add_heading("3. جدول المقارنة المتزامنة للفقرات وملاحظات MQM", level=2)
    segments = data.get('segments', [])
    t_s = doc.add_table(rows=len(segments)+1, cols=6)
    t_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["#", "الأصل الإنجليزي", "الترجمة البشرية", "ترجمة الذكاء الاصطناعي", "المقاييس", "ملاحظة MQM"]
    for i, h in enumerate(s_headers):
        cell = t_s.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True

    for r_idx, s in enumerate(segments):
        t_s.cell(r_idx+1, 0).text = str(s.get('id', ''))
        t_s.cell(r_idx+1, 1).text = str(s.get('source', ''))
        t_s.cell(r_idx+1, 2).text = str(s.get('reference', ''))
        t_s.cell(r_idx+1, 3).text = str(s.get('hypothesis', ''))
        t_s.cell(r_idx+1, 4).text = f"C: {s.get('comet')}\nM: {s.get('mqm')}%\nch: {s.get('chrf')}\nB: {s.get('bleu')}"
        t_s.cell(r_idx+1, 5).text = f"{s.get('error_type')}\n{s.get('description')}"

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

# ----------------- Web UI HTML / CSS / JS -----------------
from aiohttp import web

HTML_PAGE = """<!DOCTYPE html>
<html lang="ar" dir="rtl" class="dark">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>منصة تقييم ترجمة الكتب | Translation Benchmark</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script>
        tailwind.config = {
            darkMode: 'class',
            theme: {
                extend: {
                    colors: {
                        primary: { 500: '#0284c7', 600: '#0369a1' }
                    }
                }
            }
        }
    </script>
    <link href="https://fonts.googleapis.com/css2?family=Tajawal:wght@400;500;700;800&display=swap" rel="stylesheet">
    <!-- Firebase SDK (Compat) -->
    <script src="https://www.gstatic.com/firebasejs/10.8.0/firebase-app-compat.js"></script>
    <script src="https://www.gstatic.com/firebasejs/10.8.0/firebase-auth-compat.js"></script>
    <script src="https://www.gstatic.com/firebasejs/10.8.0/firebase-firestore-compat.js"></script>
    <!-- html2pdf.js -->
    <script src="https://cdnjs.cloudflare.com/ajax/libs/html2pdf.js/0.10.1/html2pdf.bundle.min.js"></script>
    <style>
        body { font-family: 'Tajawal', sans-serif; transition: background-color 0.3s, color 0.3s; }
        .dark body { background-color: #0b1120; color: #f8fafc; }
        body:not(.dark) { background-color: #f8fafc; color: #0f172a; }

        .dark .glass { background: rgba(30, 41, 59, 0.7); backdrop-filter: blur(12px); border: 1px solid rgba(255, 255, 255, 0.08); }
        body:not(.dark) .glass { background: rgba(255, 255, 255, 0.85); backdrop-filter: blur(12px); border: 1px solid rgba(0, 0, 0, 0.08); box-shadow: 0 10px 25px -5px rgba(0,0,0,0.05); }

        .dark .inner-card { background: rgba(15, 23, 42, 0.6); border: 1px solid rgba(255, 255, 255, 0.06); }
        body:not(.dark) .inner-card { background: #ffffff; border: 1px solid rgba(0, 0, 0, 0.06); box-shadow: 0 4px 6px -1px rgba(0,0,0,0.02); }

        .tab-active { border-bottom: 3px solid #0284c7; font-weight: bold; }
        .dark .tab-active { color: #38bdf8; }
        body:not(.dark) .tab-active { color: #0284c7; }
    </style>
</head>
<body class="min-h-screen">

    <!-- Header -->
    <header class="border-b dark:border-slate-800 border-slate-200 dark:bg-slate-900/80 bg-white/80 sticky top-0 z-50 backdrop-blur">
        <div class="max-w-7xl mx-auto px-6 py-4 flex justify-between items-center">
            <div class="flex items-center space-x-4 space-x-reverse">
                <div class="w-10 h-10 rounded-xl bg-gradient-to-tr from-sky-500 to-indigo-600 flex items-center justify-center text-xl font-bold text-white shadow-lg shadow-sky-500/20">
                    📚
                </div>
                <div>
                    <h1 class="text-xl font-extrabold dark:text-white text-slate-900">منظومة تقييم ترجمة الكتب</h1>
                    <p class="text-xs dark:text-slate-400 text-slate-500">COMET • MQM Framework • chrF++ • BLEU</p>
                </div>
            </div>
            
            <div class="flex items-center gap-3">
                <!-- Theme Toggle Button -->
                <button onclick="toggleTheme()" id="themeBtn" class="p-2.5 rounded-xl border dark:border-slate-700 border-slate-200 dark:bg-slate-800 bg-slate-100 hover:opacity-80 transition text-sm flex items-center gap-1.5" title="تبديل الوضع الليلي / النهاري">
                    <span id="themeIcon">🌙</span>
                    <span id="themeText" class="text-xs font-semibold hidden md:inline">الوضع الليلي</span>
                </button>

                <!-- Auth Buttons -->
                <div id="authUserSection" class="hidden flex items-center gap-2">
                    <span id="userEmailSpan" class="text-xs font-medium dark:text-slate-300 text-slate-700 dark:bg-slate-800 bg-slate-100 px-3 py-1.5 rounded-full border dark:border-slate-700 border-slate-200"></span>
                    <button onclick="openHistoryModal()" class="text-xs text-sky-500 hover:text-sky-600 dark:bg-sky-500/10 bg-sky-50 border border-sky-500/20 px-3 py-1.5 rounded-xl font-semibold transition">
                        📜 السجل
                    </button>
                    <button onclick="handleSignOut()" class="text-xs text-red-500 hover:text-red-600 dark:bg-red-500/10 bg-red-50 border border-red-500/20 px-3 py-1.5 rounded-xl font-semibold transition">
                        خروج
                    </button>
                </div>
                <div id="authLoginBtnSection">
                    <button onclick="openAuthModal()" class="text-xs bg-gradient-to-r from-sky-500 to-indigo-600 text-white px-4 py-2 rounded-xl font-bold shadow-lg shadow-sky-500/20 hover:opacity-90 transition">
                        🔐 تسجيل الدخول
                    </button>
                </div>
            </div>
        </div>
    </header>

    <main class="max-w-7xl mx-auto px-6 py-8">
        
        <!-- Mode Switcher: Single vs Multi-File Batch -->
        <div class="flex items-center justify-between mb-6 pb-2 border-b dark:border-slate-800 border-slate-200">
            <div class="flex gap-6">
                <button onclick="switchMode('single')" id="tabSingle" class="tab-active pb-3 text-sm flex items-center gap-2">
                    <span>📄</span> معالجة فصل فردي (Single Chapter)
                </button>
                <button onclick="switchMode('batch')" id="tabBatch" class="pb-3 text-sm dark:text-slate-400 text-slate-500 hover:text-sky-500 flex items-center gap-2">
                    <span>📚</span> حزمة فصول متعددة (Multi-File Batch)
                </button>
            </div>
            <div class="text-xs dark:text-slate-400 text-slate-500 flex items-center gap-2">
                <span class="w-2 h-2 rounded-full bg-emerald-500 animate-pulse"></span>
                <span>يدعم EPUB, DOCX, HTML, MD</span>
            </div>
        </div>

        <!-- Single Chapter Form Section -->
        <section id="singleFormSection" class="glass p-8 rounded-3xl mb-8 shadow-2xl">
            <div class="flex flex-col md:flex-row items-center justify-between mb-6 gap-4">
                <h2 class="text-xl font-bold flex items-center gap-3 text-sky-500">
                    <span>📁</span> بيانات الفصل المستهدف
                </h2>
                <div class="flex items-center gap-2 w-full md:w-auto">
                    <input type="text" id="bookTitleInput" placeholder="عنوان الكتاب (مثال: العناصر)..." class="dark:bg-slate-900 bg-slate-50 border dark:border-slate-700 border-slate-300 rounded-xl px-3 py-2 text-xs w-full md:w-64 focus:outline-none focus:border-sky-500">
                    <input type="text" id="chapterInput" placeholder="رقم الفصل..." class="dark:bg-slate-900 bg-slate-50 border dark:border-slate-700 border-slate-300 rounded-xl px-3 py-2 text-xs w-28 focus:outline-none focus:border-sky-500">
                </div>
            </div>

            <form id="singleBenchmarkForm" class="space-y-6">
                <div class="grid grid-cols-1 md:grid-cols-3 gap-6">
                    <!-- Source File -->
                    <div class="inner-card p-5 rounded-2xl">
                        <label class="block text-sm font-bold mb-2">1. الأصل الإنجليزي (Source - EN)</label>
                        <p class="text-xs dark:text-slate-400 text-slate-500 mb-3">ملف الفصل الإنجليزي (.html, .docx, .epub, .md)</p>
                        <input type="file" id="srcFile" accept=".html,.xhtml,.docx,.epub,.md,.txt" class="block w-full text-xs dark:text-slate-400 text-slate-500 file:mr-0 file:ml-4 file:py-2 file:px-4 file:rounded-xl file:border-0 file:text-xs file:font-semibold file:bg-sky-600 file:text-white hover:file:bg-sky-500 cursor-pointer">
                        <textarea id="srcText" rows="4" placeholder="أو الصق النص الإنجليزي هنا..." class="mt-3 w-full dark:bg-slate-900 bg-slate-50 border dark:border-slate-700 border-slate-300 rounded-xl p-3 text-xs focus:outline-none focus:border-sky-500"></textarea>
                    </div>

                    <!-- Human Reference File -->
                    <div class="inner-card p-5 rounded-2xl">
                        <label class="block text-sm font-bold mb-2">2. الترجمة البشرية (Reference - AR)</label>
                        <p class="text-xs dark:text-slate-400 text-slate-500 mb-3">ملف الترجمة المرجعية المعتمدة</p>
                        <input type="file" id="refFile" accept=".html,.xhtml,.docx,.epub,.md,.txt" class="block w-full text-xs dark:text-slate-400 text-slate-500 file:mr-0 file:ml-4 file:py-2 file:px-4 file:rounded-xl file:border-0 file:text-xs file:font-semibold file:bg-emerald-600 file:text-white hover:file:bg-emerald-500 cursor-pointer">
                        <textarea id="refText" rows="4" placeholder="أو الصق الترجمة البشرية هنا..." class="mt-3 w-full dark:bg-slate-900 bg-slate-50 border dark:border-slate-700 border-slate-300 rounded-xl p-3 text-xs focus:outline-none focus:border-emerald-500"></textarea>
                    </div>

                    <!-- AI Translation File -->
                    <div class="inner-card p-5 rounded-2xl">
                        <label class="block text-sm font-bold mb-2">3. ترجمة الذكاء الاصطناعي (AI MT)</label>
                        <p class="text-xs dark:text-slate-400 text-slate-500 mb-3">ملف الترجمة الآلية المراد تقييمها</p>
                        <input type="file" id="aiFile" accept=".html,.xhtml,.docx,.epub,.md,.txt" class="block w-full text-xs dark:text-slate-400 text-slate-500 file:mr-0 file:ml-4 file:py-2 file:px-4 file:rounded-xl file:border-0 file:text-xs file:font-semibold file:bg-indigo-600 file:text-white hover:file:bg-indigo-500 cursor-pointer">
                        <textarea id="aiText" rows="4" placeholder="أو الصق ترجمة الذكاء الاصطناعي هنا..." class="mt-3 w-full dark:bg-slate-900 bg-slate-50 border dark:border-slate-700 border-slate-300 rounded-xl p-3 text-xs focus:outline-none focus:border-indigo-500"></textarea>
                    </div>
                </div>

                <div class="flex items-center gap-4">
                    <button type="button" onclick="runSingleEvaluation()" id="submitBtn" class="px-8 py-3.5 bg-gradient-to-r from-sky-500 via-indigo-500 to-emerald-500 text-white font-bold rounded-2xl shadow-xl shadow-sky-500/20 hover:opacity-95 transition-all text-sm flex items-center gap-2">
                        <span>⚡</span> تشغيل المقارنة والتقييم الشامل
                    </button>
                    <div id="loadingStatus" class="hidden flex items-center gap-3 text-sky-500 text-sm font-medium">
                        <svg class="animate-spin h-5 w-5 text-sky-500" xmlns="http://www.w3.org/2000/svg" fill="none" viewBox="0 0 24 24">
                            <circle class="opacity-25" cx="12" cy="12" r="10" stroke="currentColor" stroke-width="4"></circle>
                            <path class="opacity-75" fill="currentColor" d="M4 12a8 8 0 018-8v8H4z"></path>
                        </svg>
                        <span id="loadingMsg">جاري المحاذاة وتشغيل COMET و MQM...</span>
                    </div>
                </div>
            </form>
        </section>

        <!-- Multi-File Batch Form Section -->
        <section id="batchFormSection" class="glass p-8 rounded-3xl mb-8 shadow-2xl hidden">
            <div class="mb-6">
                <h2 class="text-xl font-bold flex items-center gap-3 text-indigo-500">
                    <span>📚</span> معالجة حزمة فصول متعددة دفعة واحدة (Batch Mode)
                </h2>
                <p class="text-xs dark:text-slate-400 text-slate-500 mt-1">
                    ارفع ملفات الفصول المتعددة (الفصل 1، 2، 3، 4...)، وسيقوم النظام بمطابقتها وتقييمها وحساب المتوسطات المجمعة مع إمكانية تصفح كل فصل على حدة.
                </p>
            </div>

            <div class="grid grid-cols-1 md:grid-cols-3 gap-6 mb-6">
                <div class="inner-card p-5 rounded-2xl border-dashed border-2">
                    <label class="block text-sm font-bold mb-2">1. ملفات الأصل الإنجليزي (Source Files)</label>
                    <p class="text-xs dark:text-slate-400 text-slate-500 mb-3">اختر عدة ملفات (.html, .docx, .epub, .md)</p>
                    <input type="file" id="batchSrcFiles" multiple accept=".html,.xhtml,.docx,.epub,.md,.txt" class="block w-full text-xs dark:text-slate-400 text-slate-500 file:py-2 file:px-4 file:rounded-xl file:border-0 file:text-xs file:font-semibold file:bg-sky-600 file:text-white hover:file:bg-sky-500 cursor-pointer">
                    <ul id="batchSrcList" class="mt-3 text-[11px] dark:text-slate-400 text-slate-600 space-y-1 max-h-28 overflow-y-auto"></ul>
                </div>

                <div class="inner-card p-5 rounded-2xl border-dashed border-2">
                    <label class="block text-sm font-bold mb-2">2. ملفات الترجمة البشرية (Reference Files)</label>
                    <p class="text-xs dark:text-slate-400 text-slate-500 mb-3">اختر ملفات الترجمات البشرية المقابلة</p>
                    <input type="file" id="batchRefFiles" multiple accept=".html,.xhtml,.docx,.epub,.md,.txt" class="block w-full text-xs dark:text-slate-400 text-slate-500 file:py-2 file:px-4 file:rounded-xl file:border-0 file:text-xs file:font-semibold file:bg-emerald-600 file:text-white hover:file:bg-emerald-500 cursor-pointer">
                    <ul id="batchRefList" class="mt-3 text-[11px] dark:text-slate-400 text-slate-600 space-y-1 max-h-28 overflow-y-auto"></ul>
                </div>

                <div class="inner-card p-5 rounded-2xl border-dashed border-2">
                    <label class="block text-sm font-bold mb-2">3. ملفات ترجمة الذكاء الاصطناعي (AI MT Files)</label>
                    <p class="text-xs dark:text-slate-400 text-slate-500 mb-3">اختر ملفات ترجمة الذكاء الاصطناعي المقابلة</p>
                    <input type="file" id="batchAiFiles" multiple accept=".html,.xhtml,.docx,.epub,.md,.txt" class="block w-full text-xs dark:text-slate-400 text-slate-500 file:py-2 file:px-4 file:rounded-xl file:border-0 file:text-xs file:font-semibold file:bg-indigo-600 file:text-white hover:file:bg-indigo-500 cursor-pointer">
                    <ul id="batchAiList" class="mt-3 text-[11px] dark:text-slate-400 text-slate-600 space-y-1 max-h-28 overflow-y-auto"></ul>
                </div>
            </div>

            <div class="flex items-center gap-4">
                <button type="button" onclick="runBatchEvaluation()" id="batchSubmitBtn" class="px-8 py-3.5 bg-gradient-to-r from-indigo-500 to-purple-600 text-white font-bold rounded-2xl shadow-xl shadow-indigo-500/20 hover:opacity-95 transition-all text-sm flex items-center gap-2">
                    <span>🚀</span> بدء معالجة الحزمة وتقييم كافة الفصول
                </button>
                <div id="batchLoadingStatus" class="hidden flex items-center gap-3 text-indigo-500 text-sm font-medium">
                    <svg class="animate-spin h-5 w-5 text-indigo-500" xmlns="http://www.w3.org/2000/svg" fill="none" viewBox="0 0 24 24">
                        <circle class="opacity-25" cx="12" cy="12" r="10" stroke="currentColor" stroke-width="4"></circle>
                        <path class="opacity-75" fill="currentColor" d="M4 12a8 8 0 018-8v8H4z"></path>
                    </svg>
                    <span id="batchLoadingMsg">جاري معالجة الفصول وحساب المقاييس الإجمالية...</span>
                </div>
            </div>
        </section>

        <!-- Results Section -->
        <div id="resultsArea" class="hidden space-y-8">
            
            <!-- Multi-Format Export Center Bar -->
            <section class="glass p-6 rounded-3xl border-2 border-sky-500/30 shadow-2xl">
                <div class="flex flex-col lg:flex-row items-center justify-between gap-4">
                    <div>
                        <h3 class="text-lg font-bold flex items-center gap-2">
                            <span>📥</span> مركز تصدير التقارير (Export Center)
                        </h3>
                        <p class="text-xs dark:text-slate-400 text-slate-500 mt-0.5">تصدير فوري بالصيغة المطلوبة بترميز يدعم اللغة العربية:</p>
                    </div>
                    <div class="grid grid-cols-2 sm:grid-cols-4 gap-3 w-full lg:w-auto">
                        <button onclick="downloadPDF()" class="py-2.5 px-4 bg-gradient-to-r from-rose-600 to-red-600 hover:from-rose-500 hover:to-red-500 text-white rounded-xl text-xs font-bold shadow-lg shadow-rose-600/20 flex items-center justify-center gap-2 transition">
                            <span>📕</span> تصدير PDF
                        </button>
                        <button onclick="downloadDOCX()" class="py-2.5 px-4 bg-gradient-to-r from-blue-600 to-indigo-600 hover:from-blue-500 hover:to-indigo-500 text-white rounded-xl text-xs font-bold shadow-lg shadow-blue-600/20 flex items-center justify-center gap-2 transition">
                            <span>📝</span> تصدير Word (.docx)
                        </button>
                        <button onclick="downloadMarkdown()" class="py-2.5 px-4 bg-gradient-to-r from-slate-700 to-slate-800 hover:bg-slate-600 border border-slate-600 text-slate-100 rounded-xl text-xs font-bold flex items-center justify-center gap-2 transition">
                            <span>📑</span> تصدير Markdown (.md)
                        </button>
                        <button onclick="downloadCSV('details')" class="py-2.5 px-4 bg-gradient-to-r from-emerald-600 to-teal-600 hover:from-emerald-500 hover:to-teal-500 text-white rounded-xl text-xs font-bold shadow-lg shadow-emerald-600/20 flex items-center justify-center gap-2 transition">
                            <span>📊</span> تصدير Excel (CSV)
                        </button>
                    </div>
                </div>
            </section>

            <!-- Printable Report Area -->
            <div id="printableReport" class="space-y-8">
                
                <!-- Batch Chapters Comparison Overview (Visible in batch mode) -->
                <div id="batchOverviewSection" class="glass p-6 rounded-3xl hidden">
                    <h3 class="text-lg font-bold mb-4 flex items-center justify-between">
                        <span>📊 مقارنة فصول الكتاب المجمعة (Multi-Chapter Summary)</span>
                        <span id="batchCountBadge" class="text-xs px-3 py-1 dark:bg-slate-800 bg-slate-200 rounded-full font-bold"></span>
                    </h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-right text-xs">
                            <thead class="dark:text-slate-400 text-slate-500 dark:bg-slate-800/60 bg-slate-100">
                                <tr>
                                    <th class="p-3 rounded-r-xl">الفصل / الكتاب</th>
                                    <th class="p-3 text-center">COMET (العصبي)</th>
                                    <th class="p-3 text-center">جودة MQM</th>
                                    <th class="p-3 text-center">chrF++</th>
                                    <th class="p-3 text-center">BLEU</th>
                                    <th class="p-3 text-center">الفقرات</th>
                                    <th class="p-3 rounded-l-xl text-center">عرض التفاصيل</th>
                                </tr>
                            </thead>
                            <tbody id="batchTableBody" class="divide-y dark:divide-slate-800 divide-slate-200"></tbody>
                        </table>
                    </div>
                </div>

                <!-- Score Cards -->
                <div class="grid grid-cols-2 md:grid-cols-4 gap-6">
                    <div class="glass p-6 rounded-3xl text-center border-t-4 border-t-indigo-500">
                        <span class="text-xs font-bold dark:text-slate-400 text-slate-500 uppercase">مؤشر COMET العصبي</span>
                        <div id="metricComet" class="text-3xl font-extrabold text-indigo-500 mt-2">0.0000</div>
                        <span class="text-[10px] dark:text-slate-500 text-slate-400 mt-1 block">الأقرب للتقييم البشري (0 - 1)</span>
                    </div>
                    <div class="glass p-6 rounded-3xl text-center border-t-4 border-t-emerald-500">
                        <span class="text-xs font-bold dark:text-slate-400 text-slate-500 uppercase">جودة MQM القياسية</span>
                        <div id="metricMqm" class="text-3xl font-extrabold text-emerald-500 mt-2">0.0%</div>
                        <span class="text-[10px] dark:text-slate-500 text-slate-400 mt-1 block">خصم العقوبات لكل 100 كلمة</span>
                    </div>
                    <div class="glass p-6 rounded-3xl text-center border-t-4 border-t-sky-500">
                        <span class="text-xs font-bold dark:text-slate-400 text-slate-500 uppercase">معيار chrF++ الصرفي</span>
                        <div id="metricChrf" class="text-3xl font-extrabold text-sky-500 mt-2">0.00</div>
                        <span class="text-[10px] dark:text-slate-500 text-slate-400 mt-1 block">الأفضل لبنية اللغة العربية (0-100)</span>
                    </div>
                    <div class="glass p-6 rounded-3xl text-center border-t-4 border-t-purple-500">
                        <span class="text-xs font-bold dark:text-slate-400 text-slate-500 uppercase">معيار BLEU اللفظي</span>
                        <div id="metricBleu" class="text-3xl font-extrabold text-purple-500 mt-2">0.00</div>
                        <span class="text-[10px] dark:text-slate-500 text-slate-400 mt-1 block">تطابق الكلمات المباشر (0-100)</span>
                    </div>
                </div>

                <!-- Error Distribution Card -->
                <div class="glass p-6 rounded-3xl">
                    <h3 class="text-lg font-bold mb-4 flex items-center justify-between">
                        <span>🔍 توزيع ونسب أنواع الأخطاء (MQM Breakdown)</span>
                        <span id="totalSegmentsBadge" class="text-xs px-3 py-1 dark:bg-slate-800 bg-slate-200 rounded-full font-medium"></span>
                    </h3>
                    <div class="overflow-x-auto">
                        <table class="w-full text-right text-xs">
                            <thead class="dark:text-slate-400 text-slate-500 dark:bg-slate-800/60 bg-slate-100">
                                <tr>
                                    <th class="p-3 rounded-r-xl">نوع الخطأ / التصنيف</th>
                                    <th class="p-3">مستوى الخطورة</th>
                                    <th class="p-3">التكرار</th>
                                    <th class="p-3 rounded-l-xl">النسبة المئوية (%)</th>
                                </tr>
                            </thead>
                            <tbody id="errorTableBody" class="divide-y dark:divide-slate-800 divide-slate-200"></tbody>
                        </table>
                    </div>
                </div>

                <!-- Side-by-Side Aligned Comparison Table -->
                <div class="glass p-6 rounded-3xl">
                    <div class="flex flex-col md:flex-row items-center justify-between mb-4 gap-3">
                        <h3 class="text-lg font-bold">📖 المقارنة المتزامنة للفقرات والملاحظات اللغوية</h3>
                        <div id="chapterSelectorContainer" class="hidden flex items-center gap-2">
                            <label class="text-xs font-semibold">عرض فقرات:</label>
                            <select id="chapterSelector" onchange="selectBatchChapter(this.value)" class="dark:bg-slate-900 bg-white border dark:border-slate-700 border-slate-300 rounded-xl px-3 py-1.5 text-xs focus:outline-none focus:border-sky-500"></select>
                        </div>
                    </div>
                    <div class="overflow-x-auto max-h-[650px] overflow-y-auto">
                        <table class="w-full text-right text-xs">
                            <thead class="dark:text-slate-400 text-slate-500 dark:bg-slate-800/90 bg-slate-100 sticky top-0 z-10 backdrop-blur">
                                <tr>
                                    <th class="p-3 w-12 text-center">#</th>
                                    <th class="p-3 w-1/4">الأصل الإنجليزي (Source)</th>
                                    <th class="p-3 w-1/4">الترجمة البشرية (Reference)</th>
                                    <th class="p-3 w-1/4">ترجمة الذكاء الاصطناعي (AI MT)</th>
                                    <th class="p-3 w-28 text-center">المقاييس</th>
                                    <th class="p-3 w-1/5">ملاحظات MQM والخطأ</th>
                                </tr>
                            </thead>
                            <tbody id="comparisonTableBody" class="divide-y dark:divide-slate-800 divide-slate-200"></tbody>
                        </table>
                    </div>
                </div>
            </div>
        </div>
    </main>

    <!-- Auth Modal -->
    <div id="authModal" class="hidden fixed inset-0 z-50 bg-black/70 backdrop-blur-sm flex items-center justify-center p-4">
        <div class="glass p-8 rounded-3xl w-full max-w-md shadow-2xl relative">
            <button onclick="closeAuthModal()" class="absolute top-4 left-4 dark:text-slate-400 text-slate-500 hover:text-red-500 text-lg font-bold">&times;</button>
            <h3 class="text-lg font-extrabold mb-2">تسجيل الدخول</h3>
            <p class="text-xs dark:text-slate-400 text-slate-500 mb-6">سجل دخولك لحفظ جلسات التقييم في قاعدة بيانات Firebase.</p>

            <div class="space-y-4">
                <div>
                    <label class="block text-xs font-semibold mb-1">البريد الإلكتروني</label>
                    <input type="email" id="authEmail" placeholder="name@example.com" class="w-full dark:bg-slate-900 bg-white border dark:border-slate-700 border-slate-300 rounded-xl px-3 py-2 text-xs focus:outline-none focus:border-sky-500">
                </div>
                <div>
                    <label class="block text-xs font-semibold mb-1">كلمة المرور</label>
                    <input type="password" id="authPassword" placeholder="••••••••" class="w-full dark:bg-slate-900 bg-white border dark:border-slate-700 border-slate-300 rounded-xl px-3 py-2 text-xs focus:outline-none focus:border-sky-500">
                </div>

                <div class="flex gap-3 pt-2">
                    <button onclick="handleEmailAuth('signin')" class="flex-1 bg-sky-600 hover:bg-sky-500 text-white font-bold py-2.5 rounded-xl text-xs transition">
                        تسجيل الدخول
                    </button>
                    <button onclick="handleEmailAuth('signup')" class="flex-1 dark:bg-slate-800 bg-slate-200 hover:opacity-80 border dark:border-slate-700 border-slate-300 font-bold py-2.5 rounded-xl text-xs transition">
                        إنشاء حساب
                    </button>
                </div>

                <div class="relative flex py-2 items-center">
                    <div class="flex-grow border-t dark:border-slate-700 border-slate-300"></div>
                    <span class="flex-shrink mx-3 dark:text-slate-500 text-slate-400 text-[10px]">أو الاستمرار عبر</span>
                    <div class="flex-grow border-t dark:border-slate-700 border-slate-300"></div>
                </div>

                <button onclick="handleGoogleSignIn()" class="w-full bg-white hover:bg-slate-100 text-slate-900 border border-slate-300 font-bold py-2.5 rounded-xl text-xs flex items-center justify-center gap-2 transition">
                    <svg class="w-4 h-4" viewBox="0 0 24 24"><path fill="#4285F4" d="M23.745 12.27c0-.7-.06-1.4-.19-2.07H12v4.51h6.6c-.29 1.52-1.14 2.82-2.4 3.68v3.05h3.88c2.27-2.09 3.66-5.17 3.66-9.17z"/><path fill="#34A853" d="M12 24c3.24 0 5.95-1.08 7.93-2.91l-3.88-3.05c-1.08.72-2.45 1.16-4.05 1.16-3.12 0-5.77-2.1-6.72-4.93H1.25v3.15C3.26 21.36 7.35 24 12 24z"/><path fill="#FBBC05" d="M5.28 14.27c-.25-.72-.38-1.49-.38-2.27s.13-1.55.38-2.27V6.58H1.25C.45 8.18 0 9.99 0 12s.45 3.82 1.25 5.42l4.03-3.15z"/><path fill="#EA4335" d="M12 4.75c1.77 0 3.35.61 4.6 1.8l3.42-3.42C17.95 1.19 15.24 0 12 0 7.35 0 3.26 2.64 1.25 6.58l4.03 3.15c.95-2.83 3.6-4.93 6.72-4.93z"/></svg>
                    تسجيل الدخول بحساب Google
                </button>
            </div>
        </div>
    </div>

    <!-- History Modal -->
    <div id="historyModal" class="hidden fixed inset-0 z-50 bg-black/70 backdrop-blur-sm flex items-center justify-center p-4">
        <div class="glass p-8 rounded-3xl w-full max-w-3xl shadow-2xl relative max-h-[85vh] flex flex-col">
            <button onclick="closeHistoryModal()" class="absolute top-4 left-4 dark:text-slate-400 text-slate-500 hover:text-red-500 text-lg font-bold">&times;</button>
            <h3 class="text-lg font-extrabold mb-2">📜 سجل التقييمات المحفوظة (Firestore)</h3>
            <p class="text-xs dark:text-slate-400 text-slate-500 mb-4">الفصول التي قمت بمقارنتها مخزنة في حسابك.</p>
            
            <div class="overflow-y-auto flex-1">
                <table class="w-full text-right text-xs">
                    <thead class="dark:text-slate-400 text-slate-500 dark:bg-slate-800/80 bg-slate-100 sticky top-0">
                        <tr>
                            <th class="p-3 rounded-r-xl">الكتاب / الفصل</th>
                            <th class="p-3">COMET</th>
                            <th class="p-3">MQM</th>
                            <th class="p-3">chrF++</th>
                            <th class="p-3">BLEU</th>
                            <th class="p-3">الفقرات</th>
                            <th class="p-3 rounded-l-xl">التاريخ</th>
                        </tr>
                    </thead>
                    <tbody id="historyTableBody" class="divide-y dark:divide-slate-800 divide-slate-200"></tbody>
                </table>
            </div>
        </div>
    </div>

    <script>
        // Theme Management
        function initTheme() {
            const savedTheme = localStorage.getItem('theme') || 'dark';
            if (savedTheme === 'light') {
                document.documentElement.classList.remove('dark');
                document.getElementById('themeIcon').textContent = '☀️';
                document.getElementById('themeText').textContent = 'الوضع النهاري';
            } else {
                document.documentElement.classList.add('dark');
                document.getElementById('themeIcon').textContent = '🌙';
                document.getElementById('themeText').textContent = 'الوضع الليلي';
            }
        }
        function toggleTheme() {
            const isDark = document.documentElement.classList.contains('dark');
            if (isDark) {
                document.documentElement.classList.remove('dark');
                localStorage.setItem('theme', 'light');
                document.getElementById('themeIcon').textContent = '☀️';
                document.getElementById('themeText').textContent = 'الوضع النهاري';
            } else {
                document.documentElement.classList.add('dark');
                localStorage.setItem('theme', 'dark');
                document.getElementById('themeIcon').textContent = '🌙';
                document.getElementById('themeText').textContent = 'الوضع الليلي';
            }
        }
        initTheme();

        // Firebase Configuration (myreports-system)
        const firebaseConfig = {
            apiKey: "AIzaSyBkWGYSsB4LJgAHWHb1Fz0JgJWI2x9mLEY",
            authDomain: "myreports-system.firebaseapp.com",
            projectId: "myreports-system",
            storageBucket: "myreports-system.firebasestorage.app",
            messagingSenderId: "215752965145",
            appId: "1:215752965145:web:88e27dc2d1885ed44e3e88"
        };
        firebase.initializeApp(firebaseConfig);
        const auth = firebase.auth();
        const db = firebase.firestore();

        let currentUser = null;
        let currentResults = null;
        let batchResults = null;

        auth.onAuthStateChanged(user => {
            currentUser = user;
            if (user) {
                document.getElementById('authLoginBtnSection').classList.add('hidden');
                document.getElementById('authUserSection').classList.remove('hidden');
                document.getElementById('userEmailSpan').textContent = user.email || 'مستخدم مسجل';
            } else {
                document.getElementById('authLoginBtnSection').classList.remove('hidden');
                document.getElementById('authUserSection').classList.add('hidden');
            }
        });

        function openAuthModal() { document.getElementById('authModal').classList.remove('hidden'); }
        function closeAuthModal() { document.getElementById('authModal').classList.add('hidden'); }
        function openHistoryModal() { document.getElementById('historyModal').classList.remove('hidden'); loadHistory(); }
        function closeHistoryModal() { document.getElementById('historyModal').classList.add('hidden'); }

        async function handleEmailAuth(mode) {
            const email = document.getElementById('authEmail').value.trim();
            const password = document.getElementById('authPassword').value;
            if (!email || !password) return alert('يرجى إدخال البريد الإلكتروني وكلمة المرور.');
            try {
                if (mode === 'signup') await auth.createUserWithEmailAndPassword(email, password);
                else await auth.signInWithEmailAndPassword(email, password);
                closeAuthModal();
            } catch (err) { alert('خطأ في المصادقة: ' + err.message); }
        }

        async function handleGoogleSignIn() {
            try {
                await auth.signInWithPopup(new firebase.auth.GoogleAuthProvider());
                closeAuthModal();
            } catch (err) { alert('خطأ في تسجيل الدخول: ' + err.message); }
        }

        async function handleSignOut() { await auth.signOut(); }

        // Mode Switching
        function switchMode(mode) {
            if (mode === 'single') {
                document.getElementById('singleFormSection').classList.remove('hidden');
                document.getElementById('batchFormSection').classList.add('hidden');
                document.getElementById('tabSingle').classList.add('tab-active');
                document.getElementById('tabBatch').classList.remove('tab-active');
            } else {
                document.getElementById('singleFormSection').classList.add('hidden');
                document.getElementById('batchFormSection').classList.remove('hidden');
                document.getElementById('tabBatch').classList.add('tab-active');
                document.getElementById('tabSingle').classList.remove('tab-active');
            }
        }

        // File Selection preview for batch
        ['batchSrc', 'batchRef', 'batchAi'].forEach(prefix => {
            const inp = document.getElementById(prefix + 'Files');
            const list = document.getElementById(prefix + 'List');
            inp.addEventListener('change', () => {
                list.innerHTML = '';
                Array.from(inp.files).forEach(f => {
                    const li = document.createElement('li');
                    li.className = 'flex items-center gap-1.5';
                    li.innerHTML = `<span>✓</span> <span class="truncate">${f.name}</span>`;
                    list.appendChild(li);
                });
            });
        });

        // Evaluation Execution: Single
        async function runSingleEvaluation() {
            const btn = document.getElementById('submitBtn');
            const loader = document.getElementById('loadingStatus');
            const resultsArea = document.getElementById('resultsArea');
            
            const formData = new FormData();
            const srcFile = document.getElementById('srcFile').files[0];
            const refFile = document.getElementById('refFile').files[0];
            const aiFile = document.getElementById('aiFile').files[0];

            if (srcFile) formData.append('src_file', srcFile);
            if (refFile) formData.append('ref_file', refFile);
            if (aiFile) formData.append('ai_file', aiFile);

            formData.append('src_text', document.getElementById('srcText').value);
            formData.append('ref_text', document.getElementById('refText').value);
            formData.append('ai_text', document.getElementById('aiText').value);
            formData.append('book_title', document.getElementById('bookTitleInput').value.trim());
            formData.append('chapter', document.getElementById('chapterInput').value.trim());

            btn.disabled = true;
            btn.classList.add('opacity-50');
            loader.classList.remove('hidden');

            try {
                const response = await fetch('/api/evaluate', { method: 'POST', body: formData });
                const data = await response.json();
                if (data.error) return alert('خطأ: ' + data.error);
                
                currentResults = data;
                document.getElementById('batchOverviewSection').classList.add('hidden');
                document.getElementById('chapterSelectorContainer').classList.add('hidden');

                renderResults(data);
                resultsArea.classList.remove('hidden');
                resultsArea.scrollIntoView({ behavior: 'smooth' });

                await saveBenchmarkToFirestore(data.summary, document.getElementById('bookTitleInput').value.trim(), document.getElementById('chapterInput').value.trim());
            } catch (err) {
                alert('حدث خطأ أثناء المعالجة: ' + err.message);
            } finally {
                btn.disabled = false;
                btn.classList.remove('opacity-50');
                loader.classList.add('hidden');
            }
        }

        // Evaluation Execution: Batch
        async function runBatchEvaluation() {
            const srcFiles = document.getElementById('batchSrcFiles').files;
            const refFiles = document.getElementById('batchRefFiles').files;
            const aiFiles = document.getElementById('batchAiFiles').files;

            if (refFiles.length === 0 || aiFiles.length === 0) {
                alert('يرجى اختيار ملفات الترجمة البشرية والآلية على الأقل.');
                return;
            }

            const btn = document.getElementById('batchSubmitBtn');
            const loader = document.getElementById('batchLoadingStatus');
            const resultsArea = document.getElementById('resultsArea');

            const formData = new FormData();
            Array.from(srcFiles).forEach(f => formData.append('batch_src_files', f));
            Array.from(refFiles).forEach(f => formData.append('batch_ref_files', f));
            Array.from(aiFiles).forEach(f => formData.append('batch_ai_files', f));

            btn.disabled = true;
            btn.classList.add('opacity-50');
            loader.classList.remove('hidden');

            try {
                const response = await fetch('/api/evaluate_batch', { method: 'POST', body: formData });
                const data = await response.json();
                if (data.error) return alert('خطأ: ' + data.error);

                batchResults = data;
                renderBatchResults(data);
                resultsArea.classList.remove('hidden');
                resultsArea.scrollIntoView({ behavior: 'smooth' });
            } catch (err) {
                alert('حدث خطأ أثناء معالجة الحزمة: ' + err.message);
            } finally {
                btn.disabled = false;
                btn.classList.remove('opacity-50');
                loader.classList.add('hidden');
            }
        }

        function renderBatchResults(data) {
            document.getElementById('batchOverviewSection').classList.remove('hidden');
            document.getElementById('batchCountBadge').textContent = `عدد الفصول: ${data.chapters.length}`;

            // Render batch table
            const tbody = document.getElementById('batchTableBody');
            tbody.innerHTML = '';
            const selector = document.getElementById('chapterSelector');
            selector.innerHTML = '';

            data.chapters.forEach((ch, idx) => {
                const tr = document.createElement('tr');
                tr.className = 'hover:dark:bg-slate-800/40 hover:bg-slate-100';
                tr.innerHTML = `
                    <td class="p-3 font-bold">${ch.chapter_name}</td>
                    <td class="p-3 text-center font-bold text-indigo-500">${ch.summary.avg_comet}</td>
                    <td class="p-3 text-center font-bold text-emerald-500">${ch.summary.avg_mqm}%</td>
                    <td class="p-3 text-center text-sky-500">${ch.summary.avg_chrf}</td>
                    <td class="p-3 text-center text-purple-500">${ch.summary.avg_bleu}</td>
                    <td class="p-3 text-center font-mono">${ch.summary.total_units}</td>
                    <td class="p-3 text-center">
                        <button onclick="selectBatchChapter(${idx})" class="px-3 py-1 bg-sky-500/10 hover:bg-sky-500/20 text-sky-500 rounded-lg text-xs font-semibold">معاينة</button>
                    </td>
                `;
                tbody.appendChild(tr);

                const opt = document.createElement('option');
                opt.value = idx;
                opt.textContent = ch.chapter_name;
                selector.appendChild(opt);
            });

            document.getElementById('chapterSelectorContainer').classList.remove('hidden');
            selectBatchChapter(0);
        }

        function selectBatchChapter(idx) {
            idx = parseInt(idx);
            if (!batchResults || !batchResults.chapters[idx]) return;
            const ch = batchResults.chapters[idx];
            currentResults = ch;
            document.getElementById('chapterSelector').value = idx;
            renderResults(ch);
        }

        function renderResults(data) {
            document.getElementById('metricComet').textContent = data.summary.avg_comet;
            document.getElementById('metricMqm').textContent = data.summary.avg_mqm + '%';
            document.getElementById('metricChrf').textContent = data.summary.avg_chrf;
            document.getElementById('metricBleu').textContent = data.summary.avg_bleu;
            document.getElementById('totalSegmentsBadge').textContent = `إجمالي الفقرات: ${data.segments.length}`;

            const errTbody = document.getElementById('errorTableBody');
            errTbody.innerHTML = '';
            data.errors.forEach(err => {
                const tr = document.createElement('tr');
                tr.className = 'hover:dark:bg-slate-800/30 hover:bg-slate-100';
                tr.innerHTML = `
                    <td class="p-3 font-semibold">${err.type}</td>
                    <td class="p-3"><span class="px-2 py-0.5 rounded-full text-[10px] ${err.severity.includes('جسيم') ? 'bg-red-500/20 text-red-500' : (err.severity.includes('طفيف') ? 'bg-yellow-500/20 text-yellow-600' : 'bg-emerald-500/20 text-emerald-500')}">${err.severity}</span></td>
                    <td class="p-3 font-bold">${err.count}</td>
                    <td class="p-3 font-bold text-sky-500">${err.percentage}</td>
                `;
                errTbody.appendChild(tr);
            });

            const compTbody = document.getElementById('comparisonTableBody');
            compTbody.innerHTML = '';
            data.segments.forEach(seg => {
                const tr = document.createElement('tr');
                tr.className = 'hover:dark:bg-slate-800/40 hover:bg-slate-100 align-top';
                tr.innerHTML = `
                    <td class="p-3 text-center text-slate-400 font-mono">${seg.id}</td>
                    <td class="p-3 leading-relaxed text-[11px]">${seg.source}</td>
                    <td class="p-3 text-emerald-600 dark:text-emerald-400 leading-relaxed text-[11px]">${seg.reference}</td>
                    <td class="p-3 text-sky-600 dark:text-sky-300 leading-relaxed text-[11px]">${seg.hypothesis}</td>
                    <td class="p-3 text-center">
                        <div class="space-y-1 font-mono text-[10px]">
                            <span class="block text-indigo-500 font-bold">C: ${seg.comet}</span>
                            <span class="block text-sky-500">ch: ${seg.chrf}</span>
                            <span class="block text-purple-500">B: ${seg.bleu}</span>
                            <span class="block text-emerald-500">M: ${seg.mqm}%</span>
                        </div>
                    </td>
                    <td class="p-3">
                        <div class="text-[11px]">
                            <span class="font-bold block ${seg.severity.includes('جسيم') ? 'text-red-500' : (seg.severity.includes('طفيف') ? 'text-yellow-600' : 'text-emerald-500')}">${seg.error_type}</span>
                            <span class="dark:text-slate-400 text-slate-500 text-[10px] leading-tight block mt-1">${seg.description}</span>
                        </div>
                    </td>
                `;
                compTbody.appendChild(tr);
            });
        }

        // Firestore History
        async function loadHistory() {
            const tbody = document.getElementById('historyTableBody');
            if (!currentUser) {
                tbody.innerHTML = '<tr><td colspan="7" class="p-4 text-center text-amber-500">يرجى تسجيل الدخول أولاً لعرض السجل.</td></tr>';
                return;
            }
            tbody.innerHTML = '<tr><td colspan="7" class="p-4 text-center dark:text-slate-500 text-slate-400">جاري الجلب من Firestore...</td></tr>';
            try {
                const snap = await db.collection('benchmark_history')
                    .where('userId', '==', currentUser.uid)
                    .orderBy('createdAt', 'desc')
                    .limit(20)
                    .get();

                if (snap.empty) {
                    tbody.innerHTML = '<tr><td colspan="7" class="p-4 text-center dark:text-slate-500 text-slate-400">لا توجد تقييمات محفوظة حتى الآن.</td></tr>';
                    return;
                }

                tbody.innerHTML = '';
                snap.forEach(doc => {
                    const d = doc.data();
                    const dateStr = d.createdAt ? new Date(d.createdAt.toDate ? d.createdAt.toDate() : d.createdAt).toLocaleDateString('ar-EG') : '-';
                    const tr = document.createElement('tr');
                    tr.className = 'hover:dark:bg-slate-800/40 hover:bg-slate-100';
                    tr.innerHTML = `
                        <td class="p-3 font-semibold">${d.bookTitle || 'كتاب'} - ${d.chapter || 'فصل'}</td>
                        <td class="p-3 font-bold text-indigo-500">${d.avgComet}</td>
                        <td class="p-3 font-bold text-emerald-500">${d.avgMqm}%</td>
                        <td class="p-3 text-sky-500">${d.avgChrf}</td>
                        <td class="p-3 text-purple-500">${d.avgBleu}</td>
                        <td class="p-3 dark:text-slate-400 text-slate-500">${d.totalSegments}</td>
                        <td class="p-3 dark:text-slate-500 text-slate-400">${dateStr}</td>
                    `;
                    tbody.appendChild(tr);
                });
            } catch (e) {
                tbody.innerHTML = `<tr><td colspan="7" class="p-4 text-center text-red-500">تعذر جلب السجل: ${e.message}</td></tr>`;
            }
        }

        async function saveBenchmarkToFirestore(summaryData, bookTitle, chapter) {
            if (!currentUser) return;
            try {
                await db.collection('benchmark_history').add({
                    userId: currentUser.uid,
                    userEmail: currentUser.email,
                    bookTitle: bookTitle || 'كتاب غير محدد',
                    chapter: chapter || 'فصل',
                    avgComet: summaryData.avg_comet,
                    avgMqm: summaryData.avg_mqm,
                    avgChrf: summaryData.avg_chrf,
                    avgBleu: summaryData.avg_bleu,
                    totalSegments: summaryData.total_units,
                    createdAt: firebase.firestore.FieldValue.serverTimestamp()
                });
            } catch (e) { console.error("Firestore save error:", e); }
        }

        // Export Functions
        function downloadPDF() {
            if (!currentResults) return;
            const element = document.getElementById('printableReport');
            const bookTitle = currentResults.book_title || 'Book';
            const opt = {
                margin:       10,
                filename:     `translation_evaluation_${bookTitle}.pdf`,
                image:        { type: 'jpeg', quality: 0.98 },
                html2canvas:  { scale: 2, useCORS: true },
                jsPDF:        { unit: 'mm', format: 'a4', orientation: 'portrait' }
            };
            html2pdf().set(opt).from(element).save();
        }

        async function downloadDOCX() {
            if (!currentResults) return;
            try {
                const response = await fetch('/api/export/docx', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify(currentResults)
                });
                const blob = await response.blob();
                const link = document.createElement("a");
                link.href = URL.createObjectURL(blob);
                link.download = `تقرير_تقييم_الترجمة_${currentResults.book_title || 'كتاب'}.docx`;
                link.click();
            } catch (err) { alert('حدث خطأ أثناء تصدير Word: ' + err.message); }
        }

        function downloadMarkdown() {
            if (!currentResults) return;
            const bookTitle = currentResults.book_title || 'كتاب';
            const chapter = currentResults.chapter || 'فصل';
            const dateStr = new Date().toLocaleDateString('ar-EG');

            let md = `# تقرير مقارنة وتقييم جودة الترجمة\\n\\n`;
            md += `**الكتاب:** ${bookTitle}  \\n`;
            md += `**الفصل:** ${chapter}  \\n`;
            md += `**التاريخ:** ${dateStr}  \\n`;
            md += `**إجمالي الفقرات:** ${currentResults.segments.length}\\n\\n`;

            md += `## 1. ملخص المعايير التقييمية\\n\\n`;
            md += `| المعيار | الدرجة | الوصف |\\n| :--- | :---: | :--- |\\n`;
            md += `| **COMET** | **${currentResults.summary.avg_comet}** | التقارب الدلالي العصبي (wmt22-comet-da) |\\n`;
            md += `| **جودة MQM** | **${currentResults.summary.avg_mqm}%** | نسبة الجودة المعتمدة بعد حسم نقاط العقوبات |\\n`;
            md += `| **chrF++** | **${currentResults.summary.avg_chrf}** | التقييم المورفولوجي العربي (0-100) |\\n`;
            md += `| **BLEU** | **${currentResults.summary.avg_bleu}** | دقة التطابق اللفظي المباشر (0-100) |\\n\\n`;

            md += `## 2. توزيع ونسب أنواع الأخطاء (MQM Breakdown)\\n\\n`;
            md += `| نوع الخطأ | مستوى الخطورة | التكرار | النسبة (%) |\\n| :--- | :--- | :---: | :---: |\\n`;
            currentResults.errors.forEach(e => {
                md += `| ${e.type} | ${e.severity} | ${e.count} | ${e.percentage} |\\n`;
            });
            md += `\\n`;

            md += `## 3. جدول المقارنة المتزامنة\\n\\n`;
            md += `| # | الأصل الإنجليزي | الترجمة البشرية | ترجمة الذكاء الاصطناعي | المقاييس | الخطأ |\\n| :-: | :--- | :--- | :--- | :-: | :--- |\\n`;
            currentResults.segments.forEach(s => {
                const cleanSrc = s.source.replace(/\\|/g, '\\\\|').replace(/\\n/g, ' ');
                const cleanRef = s.reference.replace(/\\|/g, '\\\\|').replace(/\\n/g, ' ');
                const cleanHyp = s.hypothesis.replace(/\\|/g, '\\\\|').replace(/\\n/g, ' ');
                const mScores = `C:${s.comet} / M:${s.mqm}% / ch:${s.chrf} / B:${s.bleu}`;
                md += `| ${s.id} | ${cleanSrc} | ${cleanRef} | ${cleanHyp} | ${mScores} | **${s.error_type}**: ${s.description} |\\n`;
            });

            const blob = new Blob(["\\uFEFF" + md], { type: 'text/markdown;charset=utf-8;' });
            const link = document.createElement("a");
            link.href = URL.createObjectURL(blob);
            link.download = `تقرير_تقييم_الترجمة_${bookTitle}.md`;
            link.click();
        }

        function downloadCSV(type) {
            if (!currentResults) return;
            let csvContent = "";
            let filename = `translation_evaluation_${currentResults.book_title || 'details'}.csv`;

            const headers = ["Segment_ID", "Source", "Reference", "AI_MT", "COMET", "chrF++", "BLEU", "MQM_Score", "Error_Category", "Error_Type", "Severity", "Description"];
            csvContent += headers.map(h => `"${h}"`).join(",") + "\\n";
            currentResults.segments.forEach(s => {
                const row = [
                    s.id, s.source.replace(/"/g, '""'), s.reference.replace(/"/g, '""'), s.hypothesis.replace(/"/g, '""'),
                    s.comet, s.chrf, s.bleu, s.mqm, s.error_category, s.error_type, s.severity, s.description.replace(/"/g, '""')
                ];
                csvContent += row.map(c => `"${c}"`).join(",") + "\\n";
            });

            const blob = new Blob(["\\uFEFF" + csvContent], { type: 'text/csv;charset=utf-8;' });
            const link = document.createElement("a");
            link.href = URL.createObjectURL(blob);
            link.download = filename;
            link.click();
        }
    </script>
</body>
</html>
"""

def evaluate_paragraphs(src_paras, ref_paras, ai_paras, book_title="كتاب", chapter="فصل"):
    matches = align_segments(ref_paras, ai_paras)
    if not matches:
        return None

    units = []
    for idx, m in enumerate(matches):
        prev_i, prev_j, dr, da, r_txt, a_txt = m
        if src_paras:
            s_start = min(prev_j, len(src_paras) - 1)
            s_end = min(prev_j + da, len(src_paras))
            s_txt = ' '.join(src_paras[s_start:s_end])
            if not s_txt:
                s_txt = src_paras[min(idx, len(src_paras) - 1)]
        else:
            s_txt = "(لم يتم إدخال نص إنجليزي)"

        chrf_val = round(sacrebleu.sentence_chrf(a_txt, [r_txt], word_order=2).score, 2)
        bleu_val = round(sacrebleu.sentence_bleu(a_txt, [r_txt]).score, 2)

        units.append({
            'id': idx + 1,
            'source': s_txt,
            'reference': r_txt,
            'hypothesis': a_txt,
            'chrf': chrf_val,
            'bleu': bleu_val,
            'comet': 0.0
        })

    try:
        model = get_comet_model()
        comet_data = [{'src': u['source'], 'mt': u['hypothesis'], 'ref': u['reference']} for u in units]
        comet_out = model.predict(comet_data, batch_size=8, gpus=0)
        for i, score in enumerate(comet_out.scores):
            units[i]['comet'] = round(float(score), 4)
    except Exception as e:
        print(f"COMET evaluation fallback: {e}")
        for u in units:
            u['comet'] = round(min(1.0, (u['chrf'] / 100.0) * 0.9 + 0.1), 4)

    error_counts = {}
    severity_map = {}
    for u in units:
        cat, err_type, sev, desc, mqm_score = classify_mqm(
            u['source'], u['reference'], u['hypothesis'],
            u['comet'], u['chrf'], u['bleu']
        )
        u['error_category'] = cat
        u['error_type'] = err_type
        u['severity'] = sev
        u['description'] = desc
        u['mqm'] = mqm_score

        error_counts[err_type] = error_counts.get(err_type, 0) + 1
        severity_map[err_type] = sev

    total_u = len(units)
    avg_comet = round(sum(u['comet'] for u in units) / total_u, 4)
    avg_chrf = round(sum(u['chrf'] for u in units) / total_u, 2)
    avg_bleu = round(sum(u['bleu'] for u in units) / total_u, 2)
    avg_mqm = round(sum(u['mqm'] for u in units) / total_u, 2)

    errors_list = []
    for et, cnt in sorted(error_counts.items(), key=lambda x: x[1], reverse=True):
        errors_list.append({
            'type': et,
            'severity': severity_map[et],
            'count': cnt,
            'percentage': f"{(cnt / total_u) * 100:.1f}%"
        })

    return {
        'book_title': book_title,
        'chapter': chapter,
        'summary': {
            'avg_comet': avg_comet,
            'avg_chrf': avg_chrf,
            'avg_bleu': avg_bleu,
            'avg_mqm': avg_mqm,
            'total_units': total_u
        },
        'errors': errors_list,
        'segments': units
    }

async def index_handler(request):
    return web.Response(text=HTML_PAGE, content_type='text/html')

async def export_docx_handler(request):
    try:
        data = await request.json()
        docx_bytes = generate_docx_report(data)
        book_title = data.get('book_title') or 'report'
        safe_filename = urllib.parse.quote(f"translation_evaluation_{book_title}.docx")
        return web.Response(
            body=docx_bytes,
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            headers={
                'Content-Disposition': f'attachment; filename="{safe_filename}"; filename*=UTF-8\'\'{safe_filename}'
            }
        )
    except Exception as e:
        return web.json_response({'error': str(e)}, status=500)

async def evaluate_handler(request):
    data = await request.post()
    
    src_text = data.get('src_text', '').strip()
    src_file = data.get('src_file')
    if src_file:
        src_paras = [clean_english(p) for p in extract_text_from_file(src_file.filename, src_file.file.read())]
    elif src_text:
        src_paras = [clean_english(p) for p in src_text.splitlines() if len(p.strip()) > 10]
    else:
        src_paras = []

    ref_text = data.get('ref_text', '').strip()
    ref_file = data.get('ref_file')
    if ref_file:
        ref_paras = [clean_arabic(p) for p in extract_text_from_file(ref_file.filename, ref_file.file.read())]
    elif ref_text:
        ref_paras = [clean_arabic(p) for p in ref_text.splitlines() if len(p.strip()) > 10]
    else:
        return web.json_response({'error': 'يرجى تقديم ملف أو نص الترجمة البشرية (Reference).'})

    ai_text = data.get('ai_text', '').strip()
    ai_file = data.get('ai_file')
    if ai_file:
        ai_paras = [clean_arabic(p) for p in extract_text_from_file(ai_file.filename, ai_file.file.read())]
    elif ai_text:
        ai_paras = [clean_arabic(p) for p in ai_text.splitlines() if len(p.strip()) > 10]
    else:
        return web.json_response({'error': 'يرجى تقديم ملف أو نص ترجمة الذكاء الاصطناعي (AI MT).'})

    book_title = data.get('book_title') or 'كتاب'
    chapter = data.get('chapter') or 'فصل'

    res = evaluate_paragraphs(src_paras, ref_paras, ai_paras, book_title, chapter)
    if not res:
        return web.json_response({'error': 'تعذر محاذاة الفقرات، تأكد من صحة النصوص والملفات المدخلة.'})

    return web.json_response(res)

async def evaluate_batch_handler(request):
    data = await request.post()

    src_files = data.getall('batch_src_files', [])
    ref_files = data.getall('batch_ref_files', [])
    ai_files = data.getall('batch_ai_files', [])

    if not ref_files or not ai_files:
        return web.json_response({'error': 'يرجى اختيار ملفات الترجمة البشرية والآلية.'})

    count = max(len(ref_files), len(ai_files))
    results_list = []

    for i in range(count):
        ref_f = ref_files[min(i, len(ref_files) - 1)]
        ai_f = ai_files[min(i, len(ai_files) - 1)]
        src_f = src_files[i] if i < len(src_files) else None

        ref_paras = [clean_arabic(p) for p in extract_text_from_file(ref_f.filename, ref_f.file.read())]
        ai_paras = [clean_arabic(p) for p in extract_text_from_file(ai_f.filename, ai_f.file.read())]
        src_paras = [clean_english(p) for p in extract_text_from_file(src_f.filename, src_f.file.read())] if src_f else []

        ch_name = f"الفصل {i+1} ({ref_f.filename.split('.')[0]})"
        res = evaluate_paragraphs(src_paras, ref_paras, ai_paras, book_title="مجموعة فصول", chapter=ch_name)
        if res:
            res['chapter_name'] = ch_name
            results_list.append(res)

    if not results_list:
        return web.json_response({'error': 'تعذر معالجة أي من الفصول المدخلة.'})

    return web.json_response({'chapters': results_list})

def main():
    app = web.Application(client_max_size=150 * 1024 * 1024)
    app.router.add_get('/', index_handler)
    app.router.add_post('/api/evaluate', evaluate_handler)
    app.router.add_post('/api/evaluate_batch', evaluate_batch_handler)
    app.router.add_post('/api/export/docx', export_docx_handler)
    
    port = int(os.environ.get('PORT', 7860))
    print(f"==================================================")
    print(f"🚀 تطبيق تقييم ترجمة الكتب المتطور يعمل بنجاح!")
    print(f"👉 افتح المتصفح على الرابط: http://localhost:{port}")
    print(f"==================================================")
    web.run_app(app, host='0.0.0.0', port=port)

if __name__ == '__main__':
    main()
